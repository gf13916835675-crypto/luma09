"""Generate Absence_Request_September_13_Gaofeng.docx.

The original template (Absence_Request_August_1_Gaofeng.docx) was not
available in the repository, so this recreates a standard course absence
request letter with placeholders for details that only the original
template contains.
"""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Absence Request")
run.bold = True
run.font.size = Pt(16)

doc.add_paragraph("")
doc.add_paragraph("Date: September 12, 2026")
doc.add_paragraph("")
doc.add_paragraph("Dear [Teacher's Name],")
doc.add_paragraph("")
doc.add_paragraph(
    "I am writing to request a leave of absence from the [Course Name] "
    "class scheduled on Sunday, September 13, 2026. Due to [reason for "
    "absence], I will be unable to attend the class on that day."
)
doc.add_paragraph(
    "I will make sure to review the class materials and catch up on any "
    "assignments I miss. Please let me know if there is anything else I "
    "should complete in advance."
)
doc.add_paragraph(
    "Thank you very much for your understanding. I apologize for any "
    "inconvenience this may cause."
)
doc.add_paragraph("")
doc.add_paragraph("Sincerely,")
doc.add_paragraph("Gaofeng")

doc.save("Absence_Request_September_13_Gaofeng.docx")
print("saved")
