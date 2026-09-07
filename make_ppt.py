# -*- coding: utf-8 -*-
"""生成"交我算"平台一页 PPT(交大红 + 蓝配色)"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SJTU_RED = RGBColor(0xA7, 0x20, 0x38)      # 交大红
SJTU_RED_DARK = RGBColor(0x8A, 0x1A, 0x2E)
SJTU_BLUE = RGBColor(0x00, 0x45, 0x7C)     # 深蓝
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
MID_GREY = RGBColor(0x9B, 0x9B, 0x9B)
DARK_TEXT = RGBColor(0x33, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

SW = prs.slide_width
SH = prs.slide_height


def set_font(run, size, color, bold=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.color.rgb = color
    f.bold = bold
    # 中文字体需要额外设置 East Asian 字体
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', FONT)


def add_box(x, y, w, h, fill=None, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w or 1)
    return sp


def add_text(x, y, w, h, lines, anchor=MSO_ANCHOR.MIDDLE):
    """lines: list of (text, size, color, bold, align)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, (text, size, color, bold, align) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        set_font(r, size, color, bold)
    return tb

# ---------- 顶部红色标题条 ----------
add_box(0, 0, SW, Inches(1.0), fill=SJTU_RED)
add_box(0, Inches(1.0), SW, Inches(0.06), fill=SJTU_BLUE)  # 蓝色细线点缀
add_text(Inches(0.6), 0, Inches(9.5), Inches(1.0), [
    ("上海交通大学“交我算”平台  —— 强大算力,赋能科研", 28, WHITE, True, PP_ALIGN.LEFT),
])
add_text(Inches(10.2), 0, Inches(2.6), Inches(1.0), [
    ("SJTU HPC & AI", 14, RGBColor(0xF0, 0xC0, 0xC8), False, PP_ALIGN.RIGHT),
])

# ---------- 三个大数字卡片 ----------
card_y = Inches(1.35)
card_h = Inches(1.55)
card_w = Inches(3.9)
gap = Inches(0.35)
start_x = (SW - card_w * 3 - gap * 2) / 2
cards = [
    ("320", " PFLOPS", "总聚合算力", SJTU_RED),
    ("75", " PB", "聚合存储能力", SJTU_BLUE),
    ("1024", " 张", "昇腾910B 国产千卡智算", SJTU_RED),
]
for i, (num, unit, label, color) in enumerate(cards):
    x = start_x + i * (card_w + gap)
    card = add_box(x, card_y, card_w, card_h, fill=WHITE, line=color, line_w=1.5,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    card.adjustments[0] = 0.08
    add_box(x, card_y, Inches(0.12), card_h, fill=color)  # 左侧色条
    tb = slide.shapes.add_textbox(x + Inches(0.3), card_y + Inches(0.12), card_w - Inches(0.5), card_h - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run(); r1.text = num
    set_font(r1, 44, color, True)
    r2 = p1.add_run(); r2.text = unit
    set_font(r2, 20, color, True)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r3 = p2.add_run(); r3.text = label
    set_font(r3, 14, DARK_TEXT, False)

# ---------- 中部:左图片占位 + 右平台列表 ----------
mid_y = Inches(3.15)
mid_h = Inches(2.55)
img1_w = Inches(5.6)
img1 = add_box(Inches(0.6), mid_y, img1_w, mid_h, fill=LIGHT_GREY, line=MID_GREY, line_w=1)
img1.line.dash_style = None
add_text(Inches(0.6), mid_y, img1_w, mid_h, [
    ("【图片位置 1】", 16, MID_GREY, True, PP_ALIGN.CENTER),
    ("机房 / 集群实拍图", 12, MID_GREY, False, PP_ALIGN.CENTER),
])

list_x = Inches(6.5)
list_w = SW - list_x - Inches(0.6)
platforms = [
    ("致远一号", "FP16 313P · 国产千卡智算平台(2025)", SJTU_RED),
    ("思源一号", "6 PFLOPS · 高校第一 · TOP500 第132位", SJTU_BLUE),
    ("π 2.0", "2.1 PFLOPS · 26240 CPU 核", SJTU_RED),
    ("ARM 平台", "国内高校首个国产 ARM 超算", SJTU_BLUE),
]
row_h = mid_h / 4
for i, (name, desc, color) in enumerate(platforms):
    y = mid_y + int(row_h) * i
    add_box(list_x, y + Inches(0.14), Inches(0.09), Inches(0.36), fill=color)
    tb = slide.shapes.add_textbox(list_x + Inches(0.28), y, list_w - Inches(0.28), Emu(int(row_h)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r1 = p.add_run(); r1.text = name + "  "
    set_font(r1, 17, color, True)
    r2 = p.add_run(); r2.text = desc
    set_font(r2, 14, DARK_TEXT, False)

# ---------- 底部:图片占位 2 ----------
bot_y = Inches(5.95)
bot_h = Inches(1.25)
img2 = add_box(Inches(0.6), bot_y, SW - Inches(1.2), bot_h, fill=LIGHT_GREY, line=MID_GREY, line_w=1)
add_text(Inches(0.6), bot_y, SW - Inches(1.2), bot_h, [
    ("【图片位置 2】  集群网络拓扑图 / 平台架构图", 14, MID_GREY, True, PP_ALIGN.CENTER),
])

prs.save("/workspace/交我算平台介绍.pptx")
print("saved")
